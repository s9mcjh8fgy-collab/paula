"""Gera um .docx de revisão jurídica a partir da matriz.json do Contrato na Régua.

Uso: python gerar_revisao.py [saida.docx]
Mostra todas as cláusulas (inclusive as opcionais), com a condição de cada trecho,
os valores padrão das variáveis destacados e a explicação de cada cláusula.
"""
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

AQUI = Path(__file__).parent
M = json.loads((AQUI / 'matriz.json').read_text(encoding='utf-8'))

ORDINAIS = ['PRIMEIRA', 'SEGUNDA', 'TERCEIRA', 'QUARTA', 'QUINTA', 'SEXTA', 'SÉTIMA', 'OITAVA',
            'NONA', 'DÉCIMA', 'DÉCIMA PRIMEIRA', 'DÉCIMA SEGUNDA', 'DÉCIMA TERCEIRA',
            'DÉCIMA QUARTA', 'DÉCIMA QUINTA', 'DÉCIMA SEXTA', 'DÉCIMA SÉTIMA']
ROMANOS = ['I', 'II', 'III', 'IV', 'V', 'VI', 'VII', 'VIII', 'IX', 'X', 'XI', 'XII', 'XIII']
UNIDADES = ['zero', 'um', 'dois', 'três', 'quatro', 'cinco', 'seis', 'sete', 'oito', 'nove', 'dez',
            'onze', 'doze', 'treze', 'quatorze', 'quinze', 'dezesseis', 'dezessete', 'dezoito', 'dezenove']
DEZENAS = ['', '', 'vinte', 'trinta', 'quarenta', 'cinquenta', 'sessenta', 'setenta', 'oitenta', 'noventa']

NOMES = {e['id']: e['nome'] for e in M['etapas']}
NOMES.update({o['id']: f'opção “{o["rotulo"]}”' for o in M['opcoes']})

COR_EXPL = 'F3EDE3'
COR_COND = RGBColor(0x8A, 0x5A, 0x2B)


def extenso(n):
    n = int(n)
    if n < 20:
        return UNIDADES[n]
    if n < 100:
        d, u = divmod(n, 10)
        return DEZENAS[d] + ('' if u == 0 else ' e ' + UNIDADES[u])
    return str(n)


def cond_humana(expr):
    partes = []
    for termo in expr.split('||'):
        e = []
        for t in termo.split('&&'):
            t = t.strip()
            neg = t.startswith('!')
            nome = NOMES.get(t.lstrip('!'), t.lstrip('!'))
            e.append(('sem ' if neg else 'com ') + nome)
        partes.append(' e '.join(e))
    return ' ou '.join(partes)


def sombrear(par, cor):
    ppr = par._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), cor)
    ppr.append(shd)


def escrever_txt(par, txt):
    """Escreve o texto destacando variáveis e trechos condicionais."""
    pos = 0
    for m in re.finditer(r'\{\{(\?([^:}]+):(.*?)|extf?:(\w+)|(\w+))\}\}', txt):
        par.add_run(txt[pos:m.start()])
        if m.group(2):
            r = par.add_run(f'[{cond_humana(m.group(2))}: ')
            r.italic = True
            r.font.color.rgb = COR_COND
            par.add_run(m.group(3))
            r = par.add_run(']')
            r.italic = True
            r.font.color.rgb = COR_COND
        else:
            var = m.group(4) or m.group(5)
            val = M['padroes'].get(var, f'[{var}]')
            texto = f'{val} ({extenso(val)})' if m.group(4) else str(val)
            r = par.add_run(texto)
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        pos = m.end()
    par.add_run(txt[pos:])


def p(doc, txt='', bold=False, size=None, align=None, recuo=None, depois=6):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(depois)
    if align:
        par.alignment = align
    if recuo:
        par.paragraph_format.left_indent = Cm(recuo)
    if txt:
        r = par.add_run(txt)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
    return par


def tag_cond(par, quando):
    if quando:
        r = par.add_run(f'  [entra {cond_humana(quando)}]')
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = COR_COND


def explicacao(doc, txt, titulo='Por que está aqui'):
    par = p(doc, depois=8)
    sombrear(par, COR_EXPL)
    r = par.add_run(titulo + ': ')
    r.bold = True
    r.font.size = Pt(10)
    r = par.add_run(txt)
    r.font.size = Pt(10)


def item(doc, numero, it):
    par = p(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    par.add_run(f'{numero}. ').bold = True
    escrever_txt(par, it['txt'])
    tag_cond(par, it.get('quando'))
    for i, sub in enumerate(it.get('lista', [])):
        sp = p(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, recuo=1, depois=3)
        sp.add_run(f'{ROMANOS[i]} – ')
        escrever_txt(sp, sub['txt'])
        tag_cond(sp, sub.get('quando'))
    if it.get('explicacao'):
        explicacao(doc, it['explicacao'], 'Sobre este item')


def main(saida):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.5)

    p(doc, 'Contrato na Régua: matriz para revisão', bold=True, size=16)
    p(doc, f'Versão {M["versao"]} · {M["atualizado"]}. {M["fontes"]}', size=9)
    p(doc, 'Como ler', bold=True)
    leg = [
        ('Amarelo', 'valor padrão de um campo que o arquiteto preenche ou ajusta no app.'),
        ('[entra com/sem ...]', 'o trecho só aparece no contrato quando a etapa ou opção indicada estiver marcada.'),
        ('Caixa bege', 'explicação que o arquiteto vê no app ao lado da cláusula. Não vai para o Word dele.'),
    ]
    for a, b in leg:
        par = p(doc, recuo=0.5, depois=2)
        par.add_run(a + ': ').bold = True
        par.add_run(b)

    p(doc, '')
    p(doc, 'Opções que o arquiteto liga e desliga', bold=True)
    for o in M['opcoes']:
        par = p(doc, recuo=0.5, depois=2)
        par.add_run('☑ ' if o['padrao'] else '☐ ')
        par.add_run(o['rotulo'])

    doc.add_page_break()
    p(doc, 'QUADRO-RESUMO', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    quadro = ['Contratada (nome, CPF/CNPJ, CAU, endereço, representante, e-mail, telefone)',
              'Contratante (nome, CPF/CNPJ, endereço, representante, e-mail, telefone)',
              'Objeto e local do projeto, área aproximada',
              'Serviços contratados (etapas marcadas) e prazo de cada etapa em dias úteis',
              'Responsáveis pelas aprovações',
              'Canais oficiais de comunicação',
              'Honorários: valor, forma (à vista, entrada + parcelas ou por etapa), meio de pagamento',
              'Remuneração do acompanhamento ou das compras, quando houver',
              'Anexos: Anexo I (Escopo dos Serviços), proposta comercial, cronograma',
              'Observações']
    t = doc.add_table(rows=len(quadro), cols=1)
    t.style = 'Table Grid'
    for i, q in enumerate(quadro):
        t.rows[i].cells[0].text = q
    p(doc, '')

    n = 0
    for c in M['clausulas']:
        if c['titulo'] is None:
            for it in c['itens']:
                par = p(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
                escrever_txt(par, it['txt'])
            explicacao(doc, c['explicacao'])
            continue
        par = p(doc, f'CLÁUSULA {ORDINAIS[n]} – {c["titulo"]}', bold=True, depois=6)
        par.paragraph_format.space_before = Pt(10)
        tag_cond(par, c.get('quando'))
        n += 1
        explicacao(doc, c['explicacao'])
        k = 0
        for it in c['itens']:
            k += 1
            item(doc, f'{n}.{k}', it)
        if c.get('gerada') == 'etapas':
            for e in M['etapas']:
                k += 1
                par = p(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
                par.add_run(f'{n}.{k}. {e["nome"]}. ').bold = True
                par.add_run(e['definicao'])
                tag_cond(par, e['id'])
                for j, r in enumerate(e['ressalvas'], 1):
                    sp = p(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, recuo=1, depois=3)
                    sp.add_run(f'{n}.{k}.{j}. ')
                    sp.add_run(r)

    doc.add_page_break()
    p(doc, 'ANEXO I – ESCOPO DOS SERVIÇOS', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    p(doc, 'No app, o arquiteto marca e desmarca cada item e pode acrescentar itens próprios. '
           'Só entram as etapas contratadas.', size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    grupo = None
    for e in M['etapas']:
        if e['grupo'] != grupo:
            grupo = e['grupo']
            p(doc, grupo.upper(), bold=True, size=10).paragraph_format.space_before = Pt(10)
        t = doc.add_table(rows=2, cols=2)
        t.style = 'Table Grid'
        cab = t.rows[0].cells
        cab[0].merge(cab[1])
        cab[0].text = e['nome']
        cab[0].paragraphs[0].runs[0].bold = True
        for col, (rot, lista) in enumerate([('Está incluído', e['inclui']), ('Não está incluído', e['nao_inclui'])]):
            cel = t.rows[1].cells[col]
            cel.text = ''
            cel.paragraphs[0].add_run(rot).bold = True
            for x in lista or ['(nada específico)']:
                cel.add_paragraph('• ' + x)
        p(doc, '')

    doc.save(saida)
    print('ok', saida)


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else str(AQUI / 'Matriz para revisão.docx'))
