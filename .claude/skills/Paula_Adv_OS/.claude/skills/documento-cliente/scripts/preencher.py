"""Preenche uma CÓPIA de um modelo da pasta Credenciais e salva na pasta do cliente.

Uso:
    python preencher.py "<nome do modelo ou caminho>" "<destino.docx>" <dados.json> [--sobrescrever]

dados.json (UTF-8):
{
  "campos": [["Nome:", "Maria da Silva"], ["CPF:", "123.456.789-00"]],
      -> preenche, na ordem, a célula vazia ao lado de cada rótulo (rótulos repetidos,
         como "Rep. legal:" e "CPF:", são preenchidos na ordem em que aparecem)
  "substituir": {"NOME DO CLIENTE": "MARIA DA SILVA",
                 "[Descrever o objeto do contrato.]": "Ação de cobrança contra ..."},
      -> troca textos do modelo (placeholders entre colchetes, nomes de assinatura etc.)
  "data": "24 de setembro de 2026",
      -> preenche a linha "Blumenau/SC, ____ de ______________ de 20____."
  "sem_poderes_especificos": true,
      -> remove o título e o quadro de Poderes Específicos (procurações)
  "remover_linhas": ["Desenho industrial:"]
      -> apaga os parágrafos que começam com esses textos (ex: no Contrato (PJ)_INPI,
         a linha de Marca ou a de Desenho industrial que não for usada)
}

O modelo original nunca é alterado: ele é só lido, e o resultado é gravado no destino.
"""
import json, os, re, sys
from docx import Document
from docx.shared import Pt
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

CREDENCIAIS = os.path.normpath(os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    '..', '..', '..', '..', '..', '..', '..', '..', '5_Acervo', '2_Modelos Gerais', 'Credenciais'))
LINHA_DATA = '____ de ______________ de 20____'


def todas_tabelas(container):
    for t in container.tables:
        yield t
        for row in t.rows:
            for cell in row.cells:
                yield from todas_tabelas(cell)


def todos_paragrafos(doc):
    yield from doc.paragraphs
    for t in todas_tabelas(doc):
        for row in t.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def preencher_campos(doc, campos):
    usados, faltou = set(), []
    linhas = []
    for t in todas_tabelas(doc):
        for row in t.rows:
            cells = row.cells
            if len(cells) == 2 and not cells[1].tables:
                linhas.append(cells)
    for rotulo, valor in campos:
        alvo = None
        for i, (a, b) in enumerate(linhas):
            if i in usados:
                continue
            if a.text.strip().lower() == rotulo.strip().lower() and not b.text.strip():
                alvo = i; break
        if alvo is None:
            faltou.append(rotulo); continue
        usados.add(alvo)
        a, b = linhas[alvo]
        r = b.paragraphs[0].add_run(valor); r.font.size = Pt(8.5)
        # preenchido: rótulo e valor alinhados pelo topo (endereço longo quebra em 2 linhas)
        a.vertical_alignment = b.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    return faltou


def substituir(doc, trocas):
    pendentes = dict(trocas)
    for p in todos_paragrafos(doc):
        for chave in list(trocas):
            if chave not in p.text:
                continue
            feito = False
            for r in p.runs:
                if chave in r.text:
                    r.text = r.text.replace(chave, trocas[chave]); feito = True
            if not feito:
                # texto quebrado em vários runs: junta no primeiro, mantendo a formatação dele
                txt = p.text.replace(chave, trocas[chave])
                for r in p.runs[1:]:
                    r.text = ''
                p.runs[0].text = txt
            pendentes.pop(chave, None)
    return list(pendentes)


def remover_poderes_especificos(doc):
    body = doc.element.body
    for p in doc.paragraphs:
        if 'Poderes Específicos' in p.text:
            prox = p._p.getnext()
            body.remove(p._p)
            if prox is not None and prox.tag.endswith('}tbl'):
                body.remove(prox)
            return True
    return False


def remover_linhas(doc, inicios):
    feitos = set()
    for p in list(todos_paragrafos(doc)):
        t = p.text.strip()
        for ini in inicios:
            if t.startswith(ini):
                p._p.getparent().remove(p._p); feitos.add(ini)
    return [i for i in inicios if i not in feitos]


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    if len(args) != 3:
        sys.exit(__doc__)
    modelo, destino, dados_path = args
    if not modelo.lower().endswith('.docx'):
        modelo += '.docx'
    if not os.path.isabs(modelo):
        modelo = os.path.join(CREDENCIAIS, modelo)
    if not os.path.exists(modelo):
        sys.exit(f'Modelo não encontrado: {modelo}\nDisponíveis: ' + ', '.join(sorted(os.listdir(CREDENCIAIS))))
    destino = os.path.abspath(destino)
    if os.path.normcase(os.path.dirname(destino)).startswith(os.path.normcase(CREDENCIAIS)):
        sys.exit('Destino dentro da pasta Credenciais: proibido. Os modelos ficam intactos; salve na pasta do cliente.')
    if os.path.exists(destino) and '--sobrescrever' not in sys.argv:
        sys.exit(f'Já existe um arquivo em {destino}. Use outro nome ou --sobrescrever (só com autorização da Paula).')
    dados = json.load(open(dados_path, encoding='utf-8'))

    doc = Document(modelo)
    avisos = []
    if dados.get('campos'):
        f = preencher_campos(doc, dados['campos'])
        if f: avisos.append('Rótulos não encontrados/sem espaço vazio: ' + ', '.join(f))
    trocas = dict(dados.get('substituir', {}))
    if dados.get('data'):
        trocas[LINHA_DATA] = dados['data']
    if trocas:
        f = substituir(doc, trocas)
        if f: avisos.append('Textos não encontrados no modelo: ' + ' | '.join(f))
    if dados.get('sem_poderes_especificos'):
        if not remover_poderes_especificos(doc):
            avisos.append('Quadro de Poderes Específicos não encontrado.')

    if dados.get('remover_linhas'):
        f = remover_linhas(doc, dados['remover_linhas'])
        if f: avisos.append('Linhas a remover não encontradas: ' + ', '.join(f))

    os.makedirs(os.path.dirname(destino), exist_ok=True)
    doc.save(destino)
    restantes = sorted({m for p in todos_paragrafos(doc) for m in re.findall(r'\[[^\]]{2,}\]', p.text)})
    print('Salvo em:', destino)
    for a in avisos:
        print('AVISO:', a)
    if restantes:
        print('Placeholders ainda no documento:', ' | '.join(restantes))


if __name__ == '__main__':
    main()
