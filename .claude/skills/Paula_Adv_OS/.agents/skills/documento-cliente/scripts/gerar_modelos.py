"""Gera os modelos de credenciais (procurações, declarações, contratos) em .docx
com a identidade visual da Paula Corrêa Advocacia, a partir do conteúdo dos modelos do Canva."""
import os, re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
HDR_IMG = os.path.join(HERE, 'assets', 'hdr.png')   # cabeçalho do Papel Timbrado (logo)
FTR_IMG = os.path.join(HERE, 'assets', 'ftr.png')   # rodapé do Papel Timbrado (contatos)
# Uso: python gerar_modelos.py <pasta de saída>
# Gerar SEMPRE numa pasta temporária, conferir com preview.py e só então copiar pra Credenciais.
if len(sys.argv) < 2:
    sys.exit('Informe a pasta de saída (temporária). Nunca gerar direto na pasta Credenciais.')
OUT = sys.argv[1]
os.makedirs(OUT, exist_ok=True)

MARROM = '6D413E'; TERRA = 'F26F4D'; BEGE = 'F1EBDF'; TEXTO = '333333'; LINHA = 'CDBBA7'; CINZA = '7F7F7F'
F_TIT = 'Borna'; F_CORPO = 'Inter'
MARG = 2.4               # alinha com o logo e a linha do rodapé
LARG = 21.0 - 2 * MARG   # 16,2 cm úteis

J = WD_ALIGN_PARAGRAPH.JUSTIFY; L = WD_ALIGN_PARAGRAPH.LEFT
C = WD_ALIGN_PARAGRAPH.CENTER; R = WD_ALIGN_PARAGRAPH.RIGHT


# ---------------------------------------------------------------- utilitários XML
def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def set_font(run, name=F_CORPO, size=None, bold=None, color=None, italic=None, caps=None, spacing=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = el('w:rFonts'); rpr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic
    if color: run.font.color.rgb = RGBColor.from_string(color)
    if caps: run.font.all_caps = True
    if spacing is not None:
        rpr.append(el('w:spacing', **{'w:val': int(spacing * 20)}))
    return run


def fmt(p, align=None, before=0, after=0, line=None, left=None, keep=False):
    pf = p.paragraph_format
    if align is not None: p.alignment = align
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    if line: pf.line_spacing = line
    if left is not None: pf.left_indent = Cm(left)
    if keep: pf.keep_with_next = True
    return p


def para(container, runs=(), **kw):
    """runs: lista de (texto, dict de fonte) ou string simples."""
    p = container.add_paragraph()
    if isinstance(runs, str):
        runs = [(runs, {})]
    base = kw.pop('font', {})
    for txt, f in runs:
        set_font(p.add_run(txt), **{**base, **f})
    return fmt(p, **kw)


def shade(cell, color):
    tcPr = cell._element.get_or_add_tcPr()
    tcPr.append(el('w:shd', **{'w:val': 'clear', 'w:color': 'auto', 'w:fill': color}))


def cell_borders(cell, **sides):
    tcPr = cell._element.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        b = el('w:tcBorders'); tcPr.append(b)
    for side in ('top', 'left', 'bottom', 'right'):
        if side in sides:
            v = sides[side]
            if v is None:
                b.append(el(f'w:{side}', **{'w:val': 'nil'}))
            else:
                sz, color, val = v
                b.append(el(f'w:{side}', **{'w:val': val, 'w:sz': sz, 'w:space': 0, 'w:color': color}))


def cell_margins(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._element.get_or_add_tcPr()
    m = el('w:tcMar')
    for side, v in (('top', top), ('bottom', bottom), ('left', left), ('right', right)):
        if v is not None:
            m.append(el(f'w:{side}', **{'w:w': int(v * 567), 'w:type': 'dxa'}))
    tcPr.append(m)


def table(container, rows, widths, borders=None):
    t = container.add_table(rows=rows, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblPr = t._element.tblPr
    lay = el('w:tblLayout', **{'w:type': 'fixed'}); tblPr.append(lay)
    tb = el('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        v = (borders or {}).get(side)
        if v:
            tb.append(el(f'w:{side}', **{'w:val': 'single', 'w:sz': v[0], 'w:space': 0, 'w:color': v[1]}))
        else:
            tb.append(el(f'w:{side}', **{'w:val': 'nil'}))
    tblPr.append(tb)
    # remove o recuo padrão das células pra tabela ficar margem a margem
    tblPr.append(el('w:tblInd', **{'w:w': 0, 'w:type': 'dxa'}))
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    grid = t._element.find(qn('w:tblGrid'))
    for i, gc in enumerate(grid.findall(qn('w:gridCol'))):
        gc.set(qn('w:w'), str(int(widths[i] * 567)))
    return t


def clear_cell(cell):
    """Célula nova vem com 1 parágrafo vazio; devolve ele pra reaproveitar."""
    return cell.paragraphs[0]


def first_para(cell, runs=(), **kw):
    p = clear_cell(cell)
    if isinstance(runs, str):
        runs = [(runs, {})]
    base = kw.pop('font', {})
    for txt, f in runs:
        set_font(p.add_run(txt), **{**base, **f})
    return fmt(p, **kw)


def row_height(row, cm, exact=False):
    row.height = Cm(cm)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY if exact else WD_ROW_HEIGHT_RULE.AT_LEAST


def cant_split(row):
    row._tr.get_or_add_trPr().append(el('w:cantSplit'))


# ---------------------------------------------------------------- documento base
def novo_doc():
    d = Document()
    st = d.styles['Normal']
    st.font.name = F_CORPO; st.font.size = Pt(9.5)
    st.font.color.rgb = RGBColor.from_string(TEXTO)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), F_CORPO)
    rpr.append(el('w:lang', **{'w:val': 'pt-BR'}))
    st.paragraph_format.space_after = Pt(0); st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.line_spacing = 1.15
    s = d.sections[0]
    s.page_width = Cm(21.0); s.page_height = Cm(29.7)
    s.left_margin = s.right_margin = Cm(MARG)
    s.top_margin = Cm(4.1); s.bottom_margin = Cm(3.8)
    s.header_distance = Cm(0); s.footer_distance = Cm(0)
    # hifenização automática: melhora o justificado nas colunas estreitas
    dts = d.settings.element.find(qn('w:defaultTabStop'))
    dts.addnext(el('w:autoHyphenation'))
    d.core_properties.author = 'Paula Corrêa Sociedade Individual de Advocacia'
    return d


def montar_cabecalho(section, titulo):
    section.header.is_linked_to_previous = False
    h = section.header
    p = h.paragraphs[0]
    fmt(p, left=-MARG); p.paragraph_format.right_indent = Cm(-MARG)
    p.add_run().add_picture(HDR_IMG, width=Cm(21.0))
    # título posicionado ao lado do logo (quadro fixo na página, não empurra o texto)
    tp = h.add_paragraph()
    tp.alignment = R
    tp._p.get_or_add_pPr().insert(0, el('w:framePr', **{
        'w:w': 9000, 'w:hSpace': 0, 'w:wrap': 'around', 'w:vAnchor': 'page',
        'w:hAnchor': 'margin', 'w:xAlign': 'right', 'w:y': 1080}))
    set_font(tp.add_run(titulo), F_TIT, 24 if len(titulo) < 16 else 19, True, TERRA)


def montar_rodape(section):
    f = section.footer
    p = f.paragraphs[0]
    fmt(p, left=-MARG); p.paragraph_format.right_indent = Cm(-MARG)
    p.add_run().add_picture(FTR_IMG, width=Cm(21.0))


def set_cols(section, n):
    sp = section._sectPr
    cols = sp.find(qn('w:cols'))
    if cols is None:
        cols = el('w:cols'); sp.append(cols)
    cols.set(qn('w:num'), str(n)); cols.set(qn('w:space'), str(int(0.8 * 567)))


# ---------------------------------------------------------------- blocos
def titulo_secao(container, texto, align=L, before=12, after=6, icone=True):
    runs = []
    if icone:
        runs.append(('✳  ', {'name': 'Segoe UI Symbol', 'size': 11, 'color': TERRA}))
    runs.append((texto, {'name': F_TIT, 'size': 14, 'bold': True, 'color': MARROM}))
    return para(container, runs, align=align, before=before, after=after, keep=True)


ADVOGADA = [
    ('Paula Corrêa Sociedade Individual de Advocacia', True),
    ('OAB/SC 3523/2017', False),
    ('CNPJ: 27.653.659/0001-01', False),
    ('Endereço: Rua Itararé, 49, sala 02, Bairro Garcia, Blumenau/SC, CEP 89020-140', False),
    ('Advogada atuante pela sociedade:', False),
    ('Paula Fernanda Corrêa de Borba · OAB/SC 28.118', True),
]


def quadro_partes(d, rot_esq, campos, rot_dir='OUTORGADA · ADVOGADA', valores=None, alt=0.62):
    """Tabela de qualificação: cliente à esquerda (rótulo + espaço pra preencher), escritório à direita."""
    valores = valores or {}
    t = table(d, 2, [LARG / 2, LARG / 2])
    for i, txt in enumerate((rot_esq, rot_dir)):
        c = t.rows[0].cells[i]
        first_para(c, [(txt, {'size': 8, 'bold': True, 'color': MARROM, 'spacing': 1})], align=C, after=0)
        cell_borders(c, bottom=(6, TERRA, 'single'))
        cell_margins(c, bottom=0.15)
    esq, dirc = t.rows[1].cells
    cell_borders(esq, right=(4, LINHA, 'single'))
    cell_margins(esq, top=0.2, right=0.35, left=0)
    cell_margins(dirc, top=0.2, left=0.45, right=0)
    # campos do cliente numa sub-tabela: rótulo | linha pra preencher
    p0 = clear_cell(esq)
    sub = esq.add_table(rows=len(campos), cols=2)
    esq._element.remove(p0._p)
    sub_w = [2.5, LARG / 2 - 2.5 - 0.4]
    sub.autofit = False
    tblPr = sub._element.tblPr
    tblPr.append(el('w:tblLayout', **{'w:type': 'fixed'}))
    for gi, gc in enumerate(sub._element.find(qn('w:tblGrid')).findall(qn('w:gridCol'))):
        gc.set(qn('w:w'), str(int(sub_w[gi] * 567)))
    for r, campo in enumerate(campos):
        row = sub.rows[r]; row_height(row, alt)
        a, b = row.cells
        a.width = Cm(sub_w[0]); b.width = Cm(sub_w[1])
        a.vertical_alignment = b.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
        first_para(a, [(campo, {'size': 8.5, 'bold': True, 'color': MARROM})])
        vp = first_para(b, [(valores.get(campo, ''), {'size': 8.5})])
        vp._p.get_or_add_pPr().append(el('w:suppressAutoHyphens'))
        cell_borders(b, bottom=(4, LINHA, 'dotted'))
        cell_margins(a, left=0, right=0.1, bottom=0.03); cell_margins(b, left=0.1, right=0, bottom=0.03)
    esq.add_paragraph()  # Word exige parágrafo depois de tabela aninhada
    fmt(esq.paragraphs[-1], after=0).paragraph_format.line_spacing = Pt(2)
    # escritório
    first = True
    for txt, bold in ADVOGADA:
        runs = [(txt, {'size': 8.5, 'bold': bold, 'color': MARROM if bold else TEXTO})]
        if first:
            first_para(dirc, runs, after=3); first = False
        else:
            para(dirc, runs, after=3)
    return t


def assinaturas(d, blocos, data=True, before=18):
    """blocos: lista de (NOME EM DESTAQUE, legenda). Data à esquerda, assinaturas à direita."""
    if data:
        t = table(d, 1, [LARG * 0.5, LARG * 0.5])
        esq, dirc = t.rows[0].cells
        esq.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
        first_para(esq, [('Blumenau/SC, ____ de ______________ de 20____.', {'size': 9})], after=4)
        alvo = dirc
    else:
        alvo = None
    ass = []
    for i, (nome, leg) in enumerate(blocos):
        cont = alvo if alvo is not None else d
        if i == 0 and alvo is not None:
            p = first_para(alvo, [('', {})], before=before)
        else:
            p = para(cont, [('', {})], before=before)
        # linha de assinatura: borda superior do parágrafo do nome
        p2 = para(cont, [(nome, {'size': 8, 'bold': True, 'color': MARROM, 'spacing': 0.5})], align=C, before=0)
        pbdr = el('w:pBdr'); pbdr.append(el('w:top', **{'w:val': 'single', 'w:sz': 6, 'w:space': 4, 'w:color': MARROM}))
        p2._p.get_or_add_pPr().append(pbdr)
        if not alvo:
            p2.paragraph_format.left_indent = Cm(LARG * 0.5)
        if leg:
            p3 = para(cont, [(leg, {'size': 8, 'color': CINZA})], align=C)
            if not alvo:
                p3.paragraph_format.left_indent = Cm(LARG * 0.5)
        ass.append(p2)
    return ass


def fim(d):
    """Parágrafo final mínimo: o Word exige um depois de tabela e ele não pode empurrar pra outra página."""
    p = d.add_paragraph()
    p.paragraph_format.line_spacing = Pt(1)
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run(''), size=1)


def assinaturas_abaixo(d, blocos, larg=8.5, apertado=False):
    """Data em cima, assinaturas embaixo: uma centralizada, ou duas lado a lado."""
    para(d, [('Blumenau/SC, ____ de ______________ de 20____.', {'size': 10})], before=14 if apertado else 22)
    if len(blocos) == 2:
        assinaturas_lado_a_lado(d, blocos[0], blocos[1], before=30)
        return
    recuo = (LARG - larg) / 2
    for nome, leg in blocos:
        para(d, '', before=(26 if apertado else 38) if len(blocos) == 1 else 22)
        p2 = para(d, [(nome, {'size': 8, 'bold': True, 'color': MARROM, 'spacing': 0.5})], align=C)
        pbdr = el('w:pBdr'); pbdr.append(el('w:top', **{'w:val': 'single', 'w:sz': 6, 'w:space': 4, 'w:color': MARROM}))
        p2._p.get_or_add_pPr().append(pbdr)
        p2.paragraph_format.left_indent = p2.paragraph_format.right_indent = Cm(recuo)
        if leg:
            p3 = para(d, [(leg, {'size': 8, 'color': CINZA})], align=C)
            p3.paragraph_format.left_indent = p3.paragraph_format.right_indent = Cm(recuo)
        p2.paragraph_format.keep_together = True


def assinaturas_lado_a_lado(d, esq, dirc, before=26):
    """Duas colunas de assinatura (contrato: CONTRATANTE | CONTRATADA)."""
    t = table(d, 1, [LARG / 2, LARG / 2])
    for cell, (nome, leg) in zip(t.rows[0].cells, (esq, dirc)):
        cell_margins(cell, left=0.6, right=0.6)
        first_para(cell, '', before=before)
        p2 = para(cell, [(nome, {'size': 8, 'bold': True, 'color': MARROM, 'spacing': 0.5})], align=C)
        pbdr = el('w:pBdr'); pbdr.append(el('w:top', **{'w:val': 'single', 'w:sz': 6, 'w:space': 4, 'w:color': MARROM}))
        p2._p.get_or_add_pPr().append(pbdr)
        if leg:
            para(cell, [(leg, {'size': 7.5, 'color': CINZA})], align=C)
    cant_split(t.rows[0])
    return t


def caixa(d, texto, before=0, altura=1.4):
    """Caixa bege de preenchimento livre."""
    t = table(d, 1, [LARG])
    c = t.rows[0].cells[0]
    shade(c, BEGE); row_height(t.rows[0], altura)
    cell_margins(c, top=0.2, bottom=0.2, left=0.35, right=0.35)
    first_para(c, [(texto, {'size': 9, 'color': MARROM, 'italic': texto.startswith('[')})], align=J)
    return t


MARC = re.compile(r'^((?:[IVXL]+|[a-z])(?:\.\d+)?\))\s*(.*)$', re.S)


def clausula(d, texto, size=8.8, after=5, line=1.12):
    m = MARC.match(texto.strip())
    sub = False
    if m:
        marc, corpo = m.groups()
        sub = '.' in marc
        runs = [(marc + ' ', {'bold': True, 'color': MARROM}), (corpo, {})]
    else:
        runs = [(texto, {})]
    p = para(d, runs, font={'size': size}, align=J, after=after, line=line)
    if sub:
        p.paragraph_format.left_indent = Cm(0.45)
    return p


def marcador(p, recuo=0.45):
    p.paragraph_format.left_indent = Cm(recuo)
    p.paragraph_format.first_line_indent = Cm(-0.3)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(recuo))
    return p


def item_lista(d, texto, size=8.8, line=1.12):
    p = para(d, [('•	', {'color': TERRA, 'bold': True}), (texto, {})], font={'size': size}, align=J, after=2, line=line)
    return marcador(p)


# ---------------------------------------------------------------- PROCURAÇÃO
PODERES = [
    'Representar o outorgante em juízo ou fora dele, em toda e qualquer ação, judicial ou administrativa, que figure como autor, réu, assistente ou opoente, de qualquer forma interessado;',
    'Propor ações, apresentar defesas, contestações, Mandado de Segurança, Habeas Corpus, Medida Cautelar, entre outros procedimentos especiais;',
    'Confessar, transigir, desistir, negociar e conciliar, inclusive e isoladamente, para os fins dos artigos 334, § 10, e 359 do CPC;',
    'Receber e dar quitação, levantar quantias depositadas, firmar compromissos, requerer falência, apresentar memoriais, prestar as primeiras e as últimas declarações em inventário;',
    'Desistir, intentar de novo, renunciar, recorrer para qualquer instância ou Tribunal, requerer junto às repartições Federais, Estaduais e Municipais, da administração direta ou indireta;',
    'Agir em conjunto ou separadamente e substabelecer o presente mandato, com ou sem reserva de poderes.',
]

CAMPOS_PF = ['Nome:', 'CPF:', 'Estado civil:', 'Profissão:', 'Endereço:']
CAMPOS_INCAPAZ = ['Nome:', 'CPF:', 'Rep. legal:', 'CPF:', 'Estado civil:', 'Profissão:', 'Endereço:']
CAMPOS_PJ = ['Razão social:', 'CNPJ:', 'Endereço:', 'Rep. legal:', 'CPF:', 'Endereço:']


def procuracao(nome_arq, campos, blocos_ass, especificos='[Descrever os poderes específicos, se houver. Se não houver, apagar este quadro.]',
               valores=None, campos_list=None):
    d = novo_doc(); s = d.sections[0]
    montar_cabecalho(s, 'Procuração'); montar_rodape(s)
    varios = len(blocos_ass) > 1
    n = len(campos_list or campos)
    alt = 0.72 if n <= 6 else 0.66
    if campos_list:
        quadro_partes_lista(d, 'OUTORGANTE · CLIENTE', campos_list, alt=alt)
    else:
        quadro_partes(d, 'OUTORGANTE · CLIENTE', campos, valores=valores, alt=alt)
    titulo_secao(d, 'Poderes', before=20, after=8)
    t = table(d, 1, [LARG / 2, LARG / 2])
    for ci, itens in enumerate((PODERES[:3], PODERES[3:])):
        c = t.rows[0].cells[ci]
        cell_margins(c, left=0 if ci == 0 else 0.4, right=0.4 if ci == 0 else 0)
        for k, txt in enumerate(itens):
            runs = [(txt, {'size': 9.5})]
            if k == 0:
                first_para(c, runs, align=J, after=9, line=1.3)
            else:
                para(c, runs, align=J, after=9, line=1.3)
    titulo_secao(d, 'Poderes Específicos', before=14, after=8)
    extra = max(0, n - 5)  # cada campo a mais na qualificação tira espaço daqui
    if varios:
        caixa(d, especificos, altura=max(1.0, 2.0 - 0.5 * extra))
        assinaturas_abaixo(d, blocos_ass, apertado=True)
    else:
        caixa(d, especificos, altura=max(1.3, 2.3 - 0.6 * extra))
        assinaturas_abaixo(d, blocos_ass, apertado=extra > 0)
    fim(d)
    d.save(os.path.join(OUT, nome_arq))


def quadro_partes_lista(d, rot_esq, pares, rot_dir='OUTORGADA · ADVOGADA', alt=0.62):
    """Versão com valores já preenchidos (ex: Consisa)."""
    valores = {}
    campos = []
    for i, (k, v) in enumerate(pares):
        key = k if k not in valores else k + ' ' * i  # rótulos repetidos (Rep. legal, CPF)
        campos.append(key); valores[key] = v
    return quadro_partes(d, rot_esq, campos, rot_dir, valores, alt)


# ---------------------------------------------------------------- DECLARAÇÃO
DECL = ('Nos termos do art. 98 e incisos do Código de Processo Civil, declara não ter condições de arcar com as taxas ou as custas judiciais, '
        'os selos postais, as despesas com publicação na imprensa oficial, dispensando-se a publicação em outros meios, a indenização devida '
        'à testemunha que, quando empregada, receberá do empregador salário integral, como se em serviço estivesse, as despesas com a realização '
        'de exame de código genético (DNA) e de outros exames considerados essenciais, os honorários do advogado e do perito e a remuneração '
        'do intérprete ou do tradutor nomeado para apresentação de versão em português de documento redigido em língua estrangeira, o custo '
        'com a elaboração de memória de cálculo, quando exigida para instauração da execução, os depósitos previstos em lei para interposição '
        'de recurso, para propositura de ação e para a prática de outros atos processuais inerentes ao exercício da ampla defesa e do '
        'contraditório, e os emolumentos devidos a notários ou registradores em decorrência da prática de registro, averbação ou qualquer '
        'outro ato notarial necessário à efetivação de decisão judicial ou à continuidade de processo judicial no qual o benefício tenha sido '
        'concedido, sem prejuízo do próprio sustento e/ou de sua família.')
DECL2 = ('Assume, neste momento, plena e total responsabilidade, inclusive criminal, pelo inteiro teor deste pedido e das declarações '
         'nele lançadas, firmando o presente.')


def declaracao(nome_arq, campos, blocos_ass):
    d = novo_doc(); s = d.sections[0]
    montar_cabecalho(s, 'Declaração de Hipossuficiência'); montar_rodape(s)
    # qualificação do declarante em faixa única
    para(d, [('DECLARANTE', {'size': 8, 'bold': True, 'color': MARROM, 'spacing': 1})], align=C, after=4)
    t = table(d, len(campos), [3.0, LARG - 3.0], borders={'top': (6, TERRA)})
    for r, campo in enumerate(campos):
        row = t.rows[r]; row_height(row, 0.85 if len(campos) <= 5 else 0.68)
        a, b = row.cells
        a.vertical_alignment = b.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
        first_para(a, [(campo, {'size': 8.5, 'bold': True, 'color': MARROM})])
        first_para(b, '')
        cell_borders(b, bottom=(4, LINHA, 'dotted'))
        cell_margins(a, left=0, bottom=0.03); cell_margins(b, left=0.1, right=0, bottom=0.03)
    apertado = len(campos) > 5
    lh = 1.38 if apertado else 1.5
    titulo_secao(d, 'Declaração', before=18 if apertado else 22, after=8)
    para(d, [(DECL, {'size': 10})], align=J, after=10, line=lh)
    para(d, [(DECL2, {'size': 10})], align=J, after=0, line=lh)
    assinaturas_abaixo(d, blocos_ass, apertado=apertado)
    fim(d)
    d.save(os.path.join(OUT, nome_arq))


# ---------------------------------------------------------------- CONTRATOS
def quadro_resumo(d, linhas, alt=1.3, fs=9, fs_rot=8.5):
    titulo_secao(d, 'Quadro Resumo', align=C, before=20 if alt > 2 else 14, after=10 if alt > 2 else 8, icone=False)
    t = table(d, len(linhas), [3.6, LARG - 3.6],
              borders={'top': (6, TERRA), 'bottom': (6, TERRA), 'insideH': (4, LINHA)})
    for r, (rot, conteudo) in enumerate(linhas):
        row = t.rows[r]; cant_split(row); row_height(row, alt)
        a, b = row.cells
        shade(a, BEGE)
        a.vertical_alignment = b.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_margins(a, left=0.3, right=0.2, top=0.2, bottom=0.2)
        cell_margins(b, left=0.35, right=0.2, top=0.2, bottom=0.2)
        lp = first_para(a, [(rot, {'size': fs_rot, 'bold': True, 'color': MARROM, 'caps': True})])
        lp._p.get_or_add_pPr().append(el('w:suppressAutoHyphens'))
        primeiro = True
        for bloco in conteudo:
            if isinstance(bloco, tuple) and bloco[0] == 'li':
                p = marcador(para(b, [('•	', {'color': TERRA, 'bold': True}), (bloco[1], {})], font={'size': fs}, after=1))
                if primeiro:
                    b._element.remove(b.paragraphs[0]._p); primeiro = False
                continue
            runs = bloco if isinstance(bloco, list) else [(bloco, {})]
            if primeiro:
                first_para(b, runs, font={'size': fs}, align=L, after=3); primeiro = False
            else:
                para(b, runs, font={'size': fs}, align=L, after=3)
        for cp in b.paragraphs:  # sem hifenização no quadro (texto curto, quebra fica feia)
            cp._p.get_or_add_pPr().append(el('w:suppressAutoHyphens'))
    return t


def contrato(nome_arq, campos, linhas_quadro, clausulas, rotulo_cli='CONTRATANTE · CLIENTE', ass_pag1=True, folgado=False, quadro_folgado=None):
    """folgado=True: contratos curtos (PF/PJ) com mais espaço pra preencher o quadro e letra maior nas cláusulas."""
    d = novo_doc(); s = d.sections[0]
    montar_cabecalho(s, 'Contrato'); montar_rodape(s)
    # quadro_folgado: altura das linhas do Quadro Resumo (página 1 cheia) sem mexer nas Condições Gerais
    pag1 = folgado or quadro_folgado
    quadro_partes(d, rotulo_cli, campos, rot_dir='CONTRATADA · ADVOGADA', alt=0.7 if pag1 else 0.62)
    if quadro_folgado:
        quadro_resumo(d, linhas_quadro, alt=quadro_folgado, fs=10, fs_rot=9)
    elif folgado:
        # cada campo a mais na qualificação (PJ tem 6) tira altura do quadro
        quadro_resumo(d, linhas_quadro, alt=2.3 - 0.18 * max(0, len(campos) - 5), fs=10, fs_rot=9)
    else:
        quadro_resumo(d, linhas_quadro)
    if ass_pag1:
        assinaturas_lado_a_lado(d, ('CONTRATANTE', None), ('CONTRATADA', None), before=36 if pag1 else 30)
    # Condições Gerais em duas colunas, em página nova, com título próprio no cabeçalho
    s2 = d.add_section(WD_SECTION.NEW_PAGE)
    set_cols(s2, 2)
    montar_cabecalho(s2, 'Condições Gerais')
    s2.footer.is_linked_to_previous = True
    for c in clausulas:
        if isinstance(c, tuple) and c[0] == 'li':
            item_lista(d, c[1], **({'size': 9.4, 'line': 1.19} if folgado else {}))
        else:
            clausula(d, c, **({'size': 9.4, 'line': 1.19, 'after': 5} if folgado else {}))
    # fecho em coluna única
    s3 = d.add_section(WD_SECTION.CONTINUOUS)
    set_cols(s3, 1)
    para(d, [('Blumenau/SC, ____ de ______________ de 20____.', {'size': 10 if folgado else 9})], before=14 if folgado else 10, after=0, keep=True)
    assinaturas_lado_a_lado(d, ('CONTRATANTE', '[Nome / razão social do cliente]'),
                            ('CONTRATADA', 'Paula Corrêa Sociedade Individual de Advocacia'), before=34 if folgado else 30)
    fim(d)
    d.save(os.path.join(OUT, nome_arq))


COND_SIMPLES = [
    'a) O advogado CONTRATADO obriga-se a prestar seus serviços profissionais na defesa dos direitos da CONTRATANTE, apenas nos serviços descritos no objeto deste contrato;',
    'b) Em remuneração desses serviços, o CONTRATADO, e na sua ausência seus herdeiros, receberá o valor descrito na forma acordada, e o seu não pagamento importará em não cumprimento do serviço contratado;',
    'c) Após a assinatura do presente instrumento, os honorários contratuais serão devidos na integralidade, mesmo em caso de desistência ou revogação dos poderes conferidos em razão deste instrumento;',
    'd) Ao CONTRATANTE caberá o pagamento das custas processuais e despesas que forem necessárias ao andamento da ação, e o não pagamento das despesas ou honorários importará no não cumprimento do serviço solicitado, sem qualquer responsabilidade ao advogado CONTRATADO;',
    'e) O não pagamento na forma estabelecida importará na aplicação de multa de 10% do valor total devido, além da aplicação de juros de 1% ao mês, acrescidos de correção monetária;',
    'f) Caso haja condenação em honorários de sucumbência, estes valores permanecerão na sua integralidade ao advogado CONTRATADO, independentemente dos honorários contratados;',
    'g) Ao CONTRATANTE caberá, também, fornecer todos os documentos e informações que o advogado ora contratado lhe solicite, sendo dever do CONTRATANTE manter seus dados cadastrais atualizados, fornecendo sempre todos os meios de contato em que seja possível localizá-lo e/ou seus familiares;',
    'h) O CONTRATANTE desde já autoriza o recebimento de malas-diretas e cartões de apresentação, seja de forma física ou por qualquer mecanismo digital;',
    'i) É permitido ao CONTRATANTE que contate o advogado CONTRATADO por todos os meios disponíveis (pessoalmente, por telefone, e-mail e/ou aplicativos de mensagens eletrônicas), entretanto o advogado CONTRATADO somente se obriga a atender e/ou responder dentro do horário comercial, aqui estabelecido de segunda a sexta-feira, das 09h00min às 17h00min;',
    'j) O advogado, a seu critério, poderá deixar de recorrer de decisão que entenda inviável de ser modificada por procedimento recursal;',
    'k) O presente contrato não importa na promessa de êxito da demanda, frente ao critério subjetivo de interpretação das normas jurídicas, e o CONTRATANTE reconhece ter sido alertado de que a presente ação é de risco;',
    'l) As Partes reconhecem a veracidade, autenticidade, integridade, validade e eficácia deste instrumento, conforme o disposto no artigo 219 do Código Civil, em formato eletrônico e/ou assinado pelas Partes por meio de certificados eletrônicos, ainda que sejam certificados eletrônicos não emitidos pela ICP-Brasil. Em caso de assinatura deste instrumento em formato físico, será digitalizado pelo CONTRATADO e arquivado apenas em formato digital, equiparando-se a documento físico para todos os efeitos legais.',
]

QUADRO_SIMPLES = [
    ('Objeto', ['[Descrever o objeto do contrato.]']),
    ('Valor dos honorários', ['[Valor dos honorários: fixo e/ou percentual.]']),
    ('Forma de pagamento', ['[Descrever a forma de pagamento.]']),
    ('Vigência e foro', ['Este contrato entra em vigor na data de sua assinatura, elegendo a Comarca de Blumenau/SC para dirimir qualquer ação oriunda deste contrato.']),
]

QUADRO_ELABORACAO = [
    ('Objeto', ['Elaboração de documento(s) jurídico(s) personalizado(s), assim especificado(s):',
                ('li', '[Quantidade e tipo de documento, ex: 01 Contrato de Prestação de Serviços Profissionais de Arquitetura.]')]),
    ('Honorários', ['R$ [valor] ([valor por extenso]).',
                    'Forma de pagamento: [nº] parcela(s) de R$ [valor] ([valor por extenso]) cada, via boleto.',
                    'Vencimentos: [datas].']),
    ('Prazo de execução', ['Até ____ dias úteis, contados do cumprimento cumulativo de:',
                           ('li', 'assinatura do contrato de honorários;'),
                           ('li', 'fornecimento completo das informações necessárias pelo CONTRATANTE.')]),
]

COND_ELABORACAO = [
    'I) A CONTRATADA obriga-se a prestar seus serviços profissionais consistentes na elaboração de documentos personalizados, conforme especificações constantes no Quadro Resumo;',
    'II) Os serviços possuem natureza intelectual e técnica, sendo desenvolvidos com base nas informações fornecidas pelo CONTRATANTE e na legislação vigente aplicável. O resultado entregue consiste em instrumento jurídico adequado à finalidade informada, não se confundindo com assessoria jurídica contínua, consultoria permanente ou garantia de resultado em eventual utilização futura;',
    'III) Em remuneração desses serviços, a CONTRATADA, e na sua ausência seus herdeiros, receberá o valor descrito na primeira página, e o seu não pagamento importará em não cumprimento do serviço contratado e na aplicação de multa de 2% do valor total devido, além da aplicação de juros de 1% ao mês, acrescidos de correção monetária;',
    'IV) O prazo de execução terá início apenas após o cumprimento integral das condições previstas no Quadro Resumo.',
    'IV.1) Em caso de inércia do CONTRATANTE no fornecimento das informações necessárias e/ou solicitadas para a elaboração do documento, por período superior a 30 dias, a contar da assinatura deste instrumento, o contrato será automaticamente rescindido, sendo devida multa equivalente a 30% do valor total do contrato, a título de compensação pelas horas já dedicadas à análise do caso; na hipótese de já ter ocorrido pagamento parcial ou integral, o valor da multa será abatido, com devolução do eventual saldo remanescente ao CONTRATANTE;',
    'V) Após a entrega do documento pela CONTRATADA, o CONTRATANTE terá o prazo de 07 dias úteis para análise e eventual solicitação de ajustes.',
    'V.1) Cada solicitação de ajuste ou alteração realizada pelo CONTRATANTE implicará na recontagem integral do prazo de execução, conforme previsto no Quadro Resumo, a partir da data da solicitação.',
    'V.2) A ausência de manifestação do CONTRATANTE dentro do prazo estabelecido será considerada como aceite integral e definitivo do documento.',
    'V.3) Solicitações realizadas após a conclusão do serviço poderão ser consideradas novo serviço, sujeitas à análise de viabilidade, prazo e cobrança de honorários;',
    'VI) São obrigações do CONTRATANTE:',
    ('li', 'fornecer todas as informações necessárias, completas e verdadeiras;'),
    ('li', 'responder de forma tempestiva às solicitações da CONTRATADA;'),
    ('li', 'analisar o documento dentro do prazo estipulado;'),
    ('li', 'efetuar o pagamento conforme pactuado.'),
    'VII) O CONTRATANTE declara estar ciente de que a qualidade e a adequação do documento dependem diretamente das informações fornecidas. A omissão, erro ou atraso no envio de informações poderá impactar o resultado final, não podendo ser imputada responsabilidade à CONTRATADA;',
    'VIII) São obrigações da CONTRATADA:',
    ('li', 'elaborar o documento conforme o escopo contratado;'),
    ('li', 'observar a legislação vigente e as boas práticas jurídicas;'),
    ('li', 'cumprir os prazos ajustados, salvo hipóteses de suspensão;'),
    ('li', 'realizar ajustes dentro dos limites contratuais.'),
    'IX) É de responsabilidade do CONTRATANTE o fornecimento das informações e dos documentos necessários à prestação do serviço a ser desempenhado pela CONTRATADA, responsabilizando-se ainda pela correção e exatidão dessas informações e/ou documentos;',
    'X) O contrato poderá ser rescindido por qualquer das partes mediante comunicação escrita. Em caso de rescisão, não haverá devolução de valores referentes a serviços já iniciados ou concluídos, bem como será devido o valor proporcional ao trabalho executado;',
    'XI) As disposições deste contrato substituem e cancelam todas e quaisquer outras avenças ou acordos que as partes tenham eventualmente mantido antes de sua assinatura, quer escritas ou verbais, prevalecendo tão só o que neste instrumento ora se ajusta;',
    'XII) Caso o CONTRATANTE tenha interesse na contratação de serviços diversos daqueles descritos no objeto deste contrato, fica acordado que o valor dos honorários será orçado previamente e, em caso de concordância, a forma de pagamento será pactuada entre as partes;',
    'XIII) É permitido ao CONTRATANTE que contate a CONTRATADA por todos os meios disponíveis (pessoalmente, por telefone, e-mail e/ou aplicativos de mensagens eletrônicas), entretanto a CONTRATADA reserva-se o direito de atender e/ou responder apenas dentro do horário comercial, aqui estabelecido como sendo de segunda a sexta-feira, das 09h00min às 17h00min;',
    'XIV) Nos termos da Lei Geral de Proteção de Dados, a CONTRATADA está autorizada a realizar o tratamento de dados pessoais do CONTRATANTE e ostenta legítimo interesse em armazenar, acessar, avaliar, modificar, transferir e comunicar, sob qualquer forma e por tempo determinado em sua política de privacidade, todos e quaisquer documentos, contratos, e-mails, cartas e demais documentações relativas ao objeto desta contratação. Tal operação de dados é e sempre será realizada unicamente em apoio e promoção às atividades técnicas e intelectuais desenvolvidas internamente pela CONTRATADA;',
    'XV) O CONTRATANTE autoriza a CONTRATADA a utilizar sua marca, logotipo e nome comercial exclusivamente para fins de divulgação institucional da CONTRATADA, incluindo portfólios, materiais publicitários, apresentações comerciais e o site da CONTRATADA. A utilização da marca deverá observar as diretrizes de identidade visual fornecidas pelo CONTRATANTE, garantindo a manutenção de sua integridade e reputação;',
    'XVI) O documento elaborado pela CONTRATADA constitui obra intelectual protegida, sendo vedado ao CONTRATANTE, salvo autorização prévia e expressa da CONTRATADA, compartilhar, ceder, disponibilizar ou distribuir o documento, total ou parcialmente, a terceiros.',
    'XVI.1) A utilização do documento é estritamente limitada ao uso próprio do CONTRATANTE, para a finalidade específica para a qual foi contratado.',
    'XVI.2) O descumprimento desta cláusula sujeitará o CONTRATANTE ao pagamento de multa não compensatória equivalente a 10 (dez) vezes o valor do presente contrato, sem prejuízo de indenização por perdas e danos;',
    'XVII) O CONTRATANTE declara ciência de que a CONTRATADA não realiza cobranças, envio de boletos ou solicitações de pagamento por canais não oficiais, sendo considerados válidos exclusivamente os contatos informados no rodapé deste contrato, incluindo e-mail, telefone e WhatsApp; a CONTRATADA não se responsabiliza por prejuízos decorrentes de fraudes, golpes ou contatos realizados por terceiros que se passem por representantes do escritório, devendo o CONTRATANTE, em caso de dúvida, confirmar a veracidade das comunicações diretamente pelos canais oficiais;',
    'XVIII) Qualquer divergência entre as partes quanto à interpretação ou ao cumprimento de quaisquer das cláusulas do presente contrato será primeiramente resolvida amigavelmente por meio de negociação e, a não ser que seja alcançada uma solução conciliatória entre as partes no prazo de até 30 (trinta) dias do momento em que surgiu tal divergência, esta será resolvida no foro da Comarca de Blumenau/SC, por mais privilegiado que outro seja ou venha a ser;',
    'XIX) As Partes reconhecem a veracidade, autenticidade, integridade, validade e eficácia deste instrumento, conforme o disposto no artigo 219 do Código Civil, em formato eletrônico e/ou assinado pelas Partes por meio de certificados eletrônicos, ainda que sejam certificados eletrônicos não emitidos pela ICP-Brasil. Em caso de assinatura deste instrumento em formato físico, será digitalizado pela CONTRATADA e arquivado apenas em formato digital, equiparando-se a documento físico para todos os efeitos legais.',
]

QUADRO_ASSESSORIA = [
    ('Objeto', ['Assessoria jurídica consultiva e preventiva nas áreas:',
                ('li', 'Cível;'), ('li', 'Empresarial;'), ('li', 'Contratual;'), ('li', 'Consumidor;'),
                ('li', 'Trabalhista; e'), ('li', 'Propriedade Intelectual (INPI).'),
                'Não faz parte do objeto a defesa dos interesses em processos judiciais, seja como autor ou réu.']),
    ('Honorários', ['R$ [valor] ([valor por extenso]) mensais.', 'Forma de pagamento: boleto.', 'Vencimento: dia [__] de cada mês.']),
    ('Vigência e foro', ['Este contrato tem vigência de 12 meses, a iniciar na data de sua assinatura.',
                         'O presente contrato prorroga-se automaticamente por tempo indeterminado, caso nenhuma das partes manifeste o interesse de rescisão com 30 (trinta) dias de antecedência do seu término.',
                         'As partes elegem a Comarca de Blumenau/SC para dirimir qualquer ação oriunda deste contrato.']),
]

COND_ASSESSORIA = [
    'I) A CONTRATADA obriga-se a prestar seus serviços profissionais na defesa dos direitos da CONTRATANTE, apenas nos serviços descritos no objeto deste contrato;',
    'II) A assessoria jurídica preventiva e consultiva objeto deste contrato será prestada em forma de respostas a consultas/dúvidas, confecção e análise de documentos e pareceres, elaboração e revisão de documentos, reuniões virtuais e presenciais com agendamento prévio, sempre dentro dos parâmetros descritos no MANUAL DO CLIENTE, o qual fará parte integrante deste contrato;',
    'III) A prestação dos serviços objeto deste contrato abrange tão somente a matriz e as filiais existentes no momento da contratação; em caso de abertura de outras filiais, os termos deste contrato deverão ser repactuados;',
    'IV) Em remuneração desses serviços, a CONTRATADA, e na sua ausência seus herdeiros, receberá o valor descrito na primeira página, e o seu não pagamento importará em não cumprimento do serviço contratado e na aplicação de multa de 2% do valor total devido, além da aplicação de juros de 1% ao mês, acrescidos de correção monetária;',
    'V) Em caso de inadimplência por mais de 30 (trinta) dias, é facultado à CONTRATADA considerar rescindido o presente contrato, independentemente de qualquer aviso ou notificação, aplicando-se multa contratual de 20% (vinte por cento) do valor equivalente à soma de 12 (doze) mensalidades;',
    'VI) Não sendo prorrogado ou renovado o presente contrato, dar-se-ão por encerrados todos os serviços, mesmo que pendentes, com exceção daqueles em que a interrupção possa causar prejuízos irremediáveis a quaisquer das partes. Em caso de renovação automática, o valor da mensalidade será reajustado pelo índice de variação acumulada do INPC ou pelo percentual fixo de 5% (cinco por cento), aquele que for maior;',
    'VII) A CONTRATADA responsabiliza-se por todos os ônus e encargos trabalhistas e previdenciários resultantes da contratação e emprego das pessoas por ela contratadas para a realização dos serviços objeto do presente contrato, bem assim pelos excessos e omissões pelos mesmos porventura praticados, inexistindo qualquer relação trabalhista ou previdenciária entre as partes contratantes;',
    'VIII) A CONTRATADA não poderá ceder a terceiros, total ou parcialmente, os direitos e obrigações oriundos deste contrato, salvo prévia e expressa autorização da CONTRATANTE;',
    'IX) Todos os tributos, taxas e contribuições que incidem ou venham a incidir sobre o presente contrato serão de responsabilidade única e exclusiva da CONTRATADA;',
    'X) A CONTRATADA declara que está e sempre esteve comprometida com o comportamento ético e probo nas relações mantidas com entidades e órgãos públicos, em especial os do Poder Judiciário, abstendo-se de praticar condutas de corrupção ou fraudes que impliquem a concessão de vantagens, gratificações, comissões e/ou incentivos indevidos, com a finalidade de influenciar comportamentos ou decisões. Também declara que cumpre e zela para que seus parceiros comerciais também cumpram todas as leis que lhes são aplicáveis, incluindo as leis e demais normas de prevenção e combate a atos de corrupção, suborno ou lavagem de dinheiro, da mesma forma que jamais autorizou, ofertou, prometeu ou realizou o pagamento ou cessão, direta ou indiretamente, de qualquer suborno, desconto, compensação, restituição, vantagem ou qualquer outro pagamento ilícito a quaisquer agentes públicos e/ou membros ou representantes de qualquer autoridade governamental ou judicial, que pudesse resultar em qualquer violação à legislação anticorrupção (Lei nº 12.846/2013), e que as operações são conduzidas sempre em cumprimento de todas as leis relativas à coibição de atos de lavagem de dinheiro;',
    'XI) É de responsabilidade da CONTRATANTE o fornecimento das informações e dos documentos necessários à prestação do serviço a ser desempenhado pela CONTRATADA, responsabilizando-se ainda pela correção e exatidão dessas informações e/ou documentos;',
    'XII) As partes e seus prepostos deverão manter absoluto sigilo em relação a terceiros de todas e quaisquer informações ou especificações contidas neste contrato, documentos, treinamentos ou manuais que venham a receber, não podendo ser divulgadas por qualquer meio ou sob qualquer justificativa, com exceção das previstas em lei, sob pena de multa contratual igual ao equivalente do valor do contrato;',
    'XIII) O dever de confidencialidade previsto no presente instrumento permanecerá íntegro durante o prazo de vigência deste contrato, incluindo suas prorrogações, e mesmo após o término deste contrato, por um prazo adicional de 5 (cinco) anos, ficando a parte que descumprir tal obrigação sujeita à indenização da parte lesada pelas perdas e danos efetivamente suportados;',
    'XIV) Caso quaisquer das partes, em virtude de lei, de decisão judicial ou por determinação de qualquer autoridade governamental, for obrigada a divulgar quaisquer informações confidenciais, deverá comunicar imediatamente o fato à outra parte, de forma que esta adote as medidas cabíveis, inclusive judiciais, para preservar as informações confidenciais. Caso as medidas tomadas para preservar as informações confidenciais não obtenham êxito, deverá ser divulgada somente a parcela das informações confidenciais necessárias à satisfação do dever legal de divulgação;',
    'XV) Em cumprimento à Lei Geral de Proteção de Dados Pessoais (LGPD), a CONTRATADA se obriga a respeitar a privacidade da CONTRATANTE, comprometendo-se a proteger e manter em sigilo todos os dados pessoais fornecidos em função deste contrato;',
    'XVI) É de responsabilidade da CONTRATANTE a obtenção do consentimento para tratamento de dados sensíveis, nos termos do artigo 11, inciso I, da LGPD, bem como o consentimento de, pelo menos, um dos pais ou do responsável legal para tratamento de dados pessoais de crianças e adolescentes, nos termos do artigo 14, § 1º, também da LGPD;',
    'XVII) Nos termos do artigo 7º, incisos V, VI e IX, todos da LGPD, a CONTRATADA está autorizada a realizar o tratamento de dados pessoais da CONTRATANTE e ostenta legítimo interesse em armazenar, acessar, avaliar, modificar, transferir e comunicar, sob qualquer forma e por tempo determinado em sua política de privacidade, todos e quaisquer documentos, contratos, e-mails, cartas e demais documentações relativas ao objeto desta contratação. Tal operação de dados é e sempre será realizada unicamente em apoio e promoção às atividades técnicas e intelectuais desenvolvidas internamente pela CONTRATADA, em especial para fins de comprovação e defesa da regular prestação dos serviços advocatícios e o respectivo resguardo de direitos e responsabilidades, bem como visando à concepção e execução de trabalhos jurídicos idênticos ou similares aos desta contratação;',
    'XVIII) As disposições deste contrato substituem e cancelam todas e quaisquer outras avenças ou acordos que as partes tenham eventualmente mantido antes de sua assinatura, quer escritas ou verbais, prevalecendo tão só o que neste instrumento ora se ajusta;',
    'XIX) Qualquer alteração do presente contrato somente se efetuará por escrito, através de instrumento próprio, sendo que a tolerância, por qualquer das partes, com relação ao descumprimento de qualquer termo ou condição ajustado não será caracterizada como desistência em exigir o cumprimento de disposição nele contida, nem representará novação com relação à obrigação passada, presente ou futura, no tocante ao termo ou condição cujo descumprimento foi tolerado;',
    'XX) A infração de qualquer cláusula deste contrato, inclusive em caso de rescisão antes do fim dos primeiros 12 (doze) meses de vigência, sujeitará a parte infratora à multa no importe de 20% (vinte por cento) do valor equivalente à soma de 12 (doze) mensalidades;',
    'XXI) Caso a CONTRATANTE tenha interesse na contratação de serviços diversos daqueles descritos no objeto deste contrato, fica acordado que o valor dos honorários será orçado previamente e, em caso de concordância, a forma de pagamento será pactuada entre as partes;',
    'XXII) À CONTRATANTE caberá o pagamento de eventuais custas processuais e despesas que forem necessárias ao andamento de demandas judiciais ou não, e o não pagamento importará no não cumprimento do serviço solicitado, sem qualquer responsabilidade da CONTRATADA;',
    'XXIII) Em caso de demandas judiciais, a CONTRATADA, a seu critério, poderá deixar de recorrer de decisão que entenda inviável de ser modificada por procedimento recursal;',
    'XXIV) É permitido à CONTRATANTE que contate a CONTRATADA por todos os meios disponíveis (pessoalmente, por telefone, e-mail e/ou aplicativos de mensagens eletrônicas), entretanto a CONTRATADA reserva-se o direito de atender e/ou responder apenas dentro do horário comercial, aqui estabelecido como sendo de segunda a sexta-feira, das 09h00min às 17h00min;',
    'XXV) A CONTRATANTE autoriza a CONTRATADA a utilizar sua marca, logotipo e nome comercial exclusivamente para fins de divulgação institucional da CONTRATADA, incluindo portfólios, materiais publicitários, apresentações comerciais e o site da CONTRATADA. A utilização da marca deverá observar as diretrizes de identidade visual fornecidas pela CONTRATANTE, garantindo a manutenção de sua integridade e reputação;',
    'XXVI) Qualquer divergência entre as partes quanto à interpretação ou ao cumprimento de quaisquer das cláusulas do presente contrato será primeiramente resolvida amigavelmente por meio de negociação e, a não ser que seja alcançada uma solução conciliatória entre as partes no prazo de até 30 (trinta) dias do momento em que surgiu tal divergência, esta será resolvida no foro da Comarca de Blumenau/SC, por mais privilegiado que outro seja ou venha a ser;',
    'XXVII) As Partes reconhecem a veracidade, autenticidade, integridade, validade e eficácia deste instrumento, conforme o disposto no artigo 219 do Código Civil, em formato eletrônico e/ou assinado pelas Partes por meio de certificados eletrônicos, ainda que sejam certificados eletrônicos não emitidos pela ICP-Brasil. Em caso de assinatura deste instrumento em formato físico, será digitalizado pela CONTRATADA e arquivado apenas em formato digital, equiparando-se a documento físico para todos os efeitos legais.',
]


QUADRO_INPI = [
    ('Objeto', ['Acompanhamento administrativo de pedido(s) de registro junto ao Instituto Nacional da Propriedade Industrial (INPI), compreendendo:',
                ('li', 'pesquisa prévia de viabilidade ou de anterioridade;'),
                ('li', 'preparo e protocolo do pedido, incluindo, no desenho industrial, o relatório descritivo e a organização das figuras fornecidas pelo CONTRATANTE;'),
                ('li', 'cumprimento de exigências e manifestação em eventual oposição administrativa;'),
                ('li', 'acompanhamento das publicações na Revista da Propriedade Industrial (RPI) até a concessão e a expedição do certificado de registro.'),
                [('Marca: ', {'bold': True, 'color': MARROM}), ('[nome da marca]', {}),
                 ('   ·   Apresentação: ', {'bold': True, 'color': MARROM}), ('[nominativa / mista / figurativa]', {}),
                 ('   ·   Classe(s) NCL: ', {'bold': True, 'color': MARROM}), ('[classe(s)]', {})],
                [('Desenho industrial: ', {'bold': True, 'color': MARROM}), ('[produto / modelo]', {}),
                 ('   ·   Classificação Locarno: ', {'bold': True, 'color': MARROM}), ('[classe]', {})],
                'Serviço de natureza exclusivamente administrativa. Não contempla medidas judiciais nem notificações extrajudiciais (ver Condições Gerais).']),
    ('Honorários', ['R$ [valor] ([valor por extenso]).',
                    'Forma de pagamento: [nº] parcela(s) de R$ [valor] ([valor por extenso]) cada, via boleto. Vencimentos: [datas].',
                    [('Taxas do INPI (GRU): ', {'bold': True, 'color': MARROM}),
                     ('não incluídas nos honorários; de responsabilidade do CONTRATANTE, conforme a tabela vigente do INPI.', {})]]),
    ('Prazo de protocolo', ['Até ____ dias úteis, contados do cumprimento cumulativo de:',
                            ('li', 'assinatura deste contrato;'),
                            ('li', 'envio das informações e documentos necessários e pagamento da GRU de depósito.')]),
]

COND_INPI = [
    'I) A CONTRATADA obriga-se a prestar serviços de natureza exclusivamente administrativa perante o INPI, voltados ao registro da(s) marca(s) e/ou do(s) desenho(s) industrial(is) descrito(s) no Quadro Resumo, em nome do CONTRATANTE, que figurará como requerente e titular do(s) pedido(s);',
    'I.1) Cada marca em cada classe, e cada desenho industrial, corresponde a um pedido distinto no INPI, com honorários e retribuições oficiais próprios. A inclusão de novos pedidos depende de orçamento e aprovação prévios;',
    'II) Não estão incluídos no objeto deste contrato, e serão orçados à parte:',
    ('li', 'qualquer medida judicial, inclusive quando decorrente de oposição, conflito com marca de terceiro ou outra situação iniciada no âmbito administrativo que precise ser levada ao Poder Judiciário, seja o CONTRATANTE autor ou réu;'),
    ('li', 'o envio ou a resposta de notificações extrajudiciais, recebidas ou enviadas no curso do pedido de registro ou após a concessão do registro;'),
    ('li', 'recurso administrativo contra o indeferimento do pedido;'),
    ('li', 'procedimentos posteriores à concessão, como processo administrativo de nulidade, caducidade, prorrogação da vigência, pagamento de retribuições quinquenais do desenho industrial, anotações de transferência ou de alteração de nome/endereço do titular, e o monitoramento da marca ou do desenho industrial contra o uso ou registro por terceiros.'),
    'II.1) Surgindo a necessidade de qualquer dos serviços acima, a CONTRATADA apresentará orçamento prévio, e o serviço somente será iniciado após a aprovação do CONTRATANTE, com honorários e forma de pagamento pactuados entre as partes;',
    'III) As retribuições oficiais devidas ao INPI (taxas de depósito, de cumprimento de exigência, de concessão e expedição do certificado, retribuições quinquenais, entre outras) não estão incluídas nos honorários e são de responsabilidade do CONTRATANTE. As guias (GRU) serão emitidas pela CONTRATADA e enviadas ao CONTRATANTE para pagamento no prazo indicado.',
    'III.1) Caso a CONTRATADA efetue o pagamento de alguma guia em nome do CONTRATANTE, este deverá reembolsá-la em até 05 (cinco) dias úteis, mediante apresentação do comprovante.',
    'III.2) O não pagamento das retribuições oficiais nos prazos fixados pelo INPI pode levar ao arquivamento do pedido ou à perda do direito, sem qualquer responsabilidade da CONTRATADA;',
    'IV) O CONTRATANTE declara ciência de que a concessão do registro depende exclusivamente da análise do INPI e de eventuais manifestações de terceiros. A pesquisa prévia de viabilidade reduz riscos, mas não garante o resultado. A obrigação da CONTRATADA é de meio, e o indeferimento, a oposição de terceiros ou o arquivamento do pedido não geram devolução dos honorários;',
    'V) O prazo de análise do pedido é definido pelo INPI e pode se estender por meses ou anos. A CONTRATADA acompanhará as publicações na RPI e informará o CONTRATANTE sobre os despachos que exijam providência ou decisão;',
    'VI) São obrigações do CONTRATANTE:',
    ('li', 'fornecer informações completas e verdadeiras sobre a marca ou o desenho industrial, a atividade exercida e os produtos ou serviços a que se destina;'),
    ('li', 'enviar os documentos solicitados, como procuração, documentos pessoais ou societários, o arquivo do logotipo (marca mista ou figurativa) e as imagens ou renderizações do objeto nos padrões exigidos pelo INPI (desenho industrial);'),
    ('li', 'pagar as guias do INPI nos prazos indicados;'),
    ('li', 'responder de forma tempestiva às solicitações da CONTRATADA;'),
    ('li', 'comunicar imediatamente à CONTRATADA qualquer notificação, correspondência ou contato recebido a respeito da marca ou do desenho industrial.'),
    'VII) Em remuneração desses serviços, a CONTRATADA, e na sua ausência seus herdeiros, receberá o valor descrito na primeira página, e o seu não pagamento importará em não cumprimento do serviço contratado e na aplicação de multa de 2% do valor total devido, além da aplicação de juros de 1% ao mês, acrescidos de correção monetária;',
    'VIII) O prazo de protocolo terá início apenas após o cumprimento integral das condições previstas no Quadro Resumo. Em caso de inércia do CONTRATANTE no fornecimento das informações e documentos necessários por período superior a 30 dias, a contar da assinatura deste instrumento, o contrato será automaticamente rescindido, sendo devida multa equivalente a 30% do valor total do contrato, a título de compensação pelas horas já dedicadas à análise do caso; na hipótese de já ter ocorrido pagamento parcial ou integral, o valor da multa será abatido, com devolução do eventual saldo remanescente ao CONTRATANTE;',
    'IX) Este contrato se encerra, em relação a cada pedido, com a expedição do certificado de registro ou com a decisão definitiva do INPI que encerre o pedido, o que ocorrer primeiro;',
    'X) O contrato poderá ser rescindido por qualquer das partes mediante comunicação escrita. Em caso de rescisão, não haverá devolução de valores referentes a serviços já iniciados ou concluídos, bem como será devido o valor proporcional ao trabalho executado. Com a rescisão, a CONTRATADA deixará de acompanhar o pedido, cabendo ao CONTRATANTE constituir novo procurador perante o INPI;',
    'XI) O CONTRATANTE declara ciência de que é comum o envio de boletos, cobranças e comunicados falsos em nome do INPI ou de empresas que oferecem supostos serviços de publicação ou registro de marcas e desenhos industriais. Somente devem ser pagas as guias enviadas pela CONTRATADA pelos canais oficiais informados no rodapé deste contrato, e a CONTRATADA não se responsabiliza por pagamentos feitos a terceiros. Em caso de dúvida, o CONTRATANTE deve confirmar a veracidade da cobrança diretamente com a CONTRATADA;',
    'XII) As disposições deste contrato substituem e cancelam todas e quaisquer outras avenças ou acordos que as partes tenham eventualmente mantido antes de sua assinatura, quer escritas ou verbais, prevalecendo tão só o que neste instrumento ora se ajusta;',
    'XIII) É permitido ao CONTRATANTE que contate a CONTRATADA por todos os meios disponíveis (pessoalmente, por telefone, e-mail e/ou aplicativos de mensagens eletrônicas), entretanto a CONTRATADA reserva-se o direito de atender e/ou responder apenas dentro do horário comercial, aqui estabelecido como sendo de segunda a sexta-feira, das 09h00min às 17h00min;',
    'XIV) Nos termos da Lei Geral de Proteção de Dados, a CONTRATADA está autorizada a realizar o tratamento de dados pessoais do CONTRATANTE e ostenta legítimo interesse em armazenar, acessar, avaliar, modificar, transferir e comunicar, sob qualquer forma e por tempo determinado em sua política de privacidade, todos e quaisquer documentos, contratos, e-mails, cartas e demais documentações relativas ao objeto desta contratação. Tal operação de dados é e sempre será realizada unicamente em apoio e promoção às atividades técnicas e intelectuais desenvolvidas internamente pela CONTRATADA;',
    'XV) O CONTRATANTE autoriza a CONTRATADA a utilizar sua marca, logotipo e nome comercial exclusivamente para fins de divulgação institucional da CONTRATADA, incluindo portfólios, materiais publicitários, apresentações comerciais e o site da CONTRATADA, observadas as diretrizes de identidade visual fornecidas pelo CONTRATANTE;',
    'XVI) Qualquer divergência entre as partes quanto à interpretação ou ao cumprimento de quaisquer das cláusulas do presente contrato será primeiramente resolvida amigavelmente por meio de negociação e, a não ser que seja alcançada uma solução conciliatória entre as partes no prazo de até 30 (trinta) dias do momento em que surgiu tal divergência, esta será resolvida no foro da Comarca de Blumenau/SC, por mais privilegiado que outro seja ou venha a ser;',
    'XVII) As Partes reconhecem a veracidade, autenticidade, integridade, validade e eficácia deste instrumento, conforme o disposto no artigo 219 do Código Civil, em formato eletrônico e/ou assinado pelas Partes por meio de certificados eletrônicos, ainda que sejam certificados eletrônicos não emitidos pela ICP-Brasil. Em caso de assinatura deste instrumento em formato físico, será digitalizado pela CONTRATADA e arquivado apenas em formato digital, equiparando-se a documento físico para todos os efeitos legais.',
]


# ---------------------------------------------------------------- geração
if __name__ == '__main__':
    procuracao('Procuração (PF).docx', CAMPOS_PF, [('NOME DO CLIENTE', 'Outorgante')])
    procuracao('Procuração (PJ).docx', CAMPOS_PJ, [('RAZÃO SOCIAL DA EMPRESA', 'Representante legal: [nome]')])
    procuracao('Procuração (Relativamente Incapaz).docx', CAMPOS_INCAPAZ,
               [('NOME DO CLIENTE', 'Outorgante (relativamente incapaz)'), ('NOME DO REPRESENTANTE LEGAL', 'Assistente / representante legal')])
    procuracao('Procuração (Absolutamente Incapaz).docx', CAMPOS_INCAPAZ,
               [('NOME DO REPRESENTANTE LEGAL', 'Representante legal de [nome do cliente]')])
    procuracao('Procuração Consisa.docx', None,
               [('JOSÉ ANTONIO ANTUNES ZANROSSO', 'p/ Consisanet Sistemas de Informação Ltda'),
                ('LUCAS BIANCO ANTUNES ZANROSSO', 'p/ Consisanet Sistemas de Informação Ltda')],
               especificos='Manifestar-se no processo nº 0016151-38.2018.8.16.0083.',
               campos_list=[('Razão social:', 'Consisanet Sistemas de Informação Ltda'), ('CNPJ:', '07.784.629/0001-19'),
                            ('Endereço:', 'Rua Palmas, 1451, sala 202, Centro, Francisco Beltrão/PR, CEP [conferir]'),
                            ('Rep. legal:', 'José Antonio Antunes Zanrosso'), ('CPF:', '422.755.049-15'),
                            ('Rep. legal:', 'Lucas Bianco Antunes Zanrosso'), ('CPF:', '053.386.989-78')])
    declaracao('Declaração (PF).docx', CAMPOS_PF, [('NOME DO CLIENTE', 'Declarante')])
    declaracao('Declaração (Relativamente Incapaz).docx', CAMPOS_INCAPAZ,
               [('NOME DO CLIENTE', 'Declarante (relativamente incapaz)'), ('NOME DO REPRESENTANTE LEGAL', 'Assistente / representante legal')])
    declaracao('Declaração (Absolutamente Incapaz).docx', CAMPOS_INCAPAZ,
               [('NOME DO REPRESENTANTE LEGAL', 'Representante legal de [nome do cliente]')])
    contrato('Contrato (PF).docx', CAMPOS_PF, QUADRO_SIMPLES, COND_SIMPLES, folgado=True)
    contrato('Contrato (PJ).docx', CAMPOS_PJ, QUADRO_SIMPLES, COND_SIMPLES, folgado=True)
    contrato('Contrato (PJ)_Elaboração de Documentos.docx', CAMPOS_PJ[:5], QUADRO_ELABORACAO, COND_ELABORACAO, quadro_folgado=3.5)
    contrato('Contrato (PJ)_Assessoria Jurídica.docx', CAMPOS_PJ[:5], QUADRO_ASSESSORIA, COND_ASSESSORIA)
    contrato('Contrato (PJ)_INPI.docx', CAMPOS_PJ[:5], QUADRO_INPI, COND_INPI)
    contrato('Contrato (PF)_INPI.docx', CAMPOS_PF, QUADRO_INPI, COND_INPI)
    print('ok', OUT)
