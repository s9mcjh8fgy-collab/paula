"""Preenche uma CÓPIA de um modelo da pasta Credenciais e salva na pasta do cliente.

Uso:
    python preencher.py "<nome do modelo ou caminho>" "<destino.docx>" <dados.json> [--sobrescrever]

dados.json (UTF-8):
{
  "campos": [["Nome:", "Maria da Silva"], ["CPF:", "123.456.789-00"]],
      -> preenche, na ordem, a célula vazia ao lado de cada rótulo (rótulos repetidos,
         como "Rep. legal:" e "CPF:", são preenchidos na ordem em que aparecem)
  "substituir": {"NOME DO CLIENTE": "MARIA DA SILVA",
                 "[Descrever o objeto do contrato.]": "Ação de cobrança contra ..."},
      -> troca textos do modelo (placeholders entre colchetes, nomes de assinatura etc.)
  "data": "24 de setembro de 2026",
      -> preenche a linha "Blumenau/SC, ____ de ______________ de 20____."
  "sem_poderes_especificos": true,
      -> remove o título e o quadro de Poderes Específicos (procurações)
  "expandir_lista": {"[Quantidade e tipo de documento, ex: ...]": ["01 Termo de uso;", "01 Contrato ..."]},
      -> troca um item de lista (marcador) por vários itens, um por linha, com a mesma formatação
         (a chave pode ser só o começo do texto do item)
  "altura_quadro": 2.0,
      -> altura mínima (cm) das linhas do Quadro Resumo; usar quando o conteúdo preenchido já é
         grande e as assinaturas da página 1 foram empurradas pra página seguinte
  "remover_linhas": ["Desenho industrial:"]
      -> apaga os parágrafos que começam com esses textos (ex: no Contrato (PJ)_INPI,
         a linha de Marca ou a de Desenho industrial que não for usada)
}

O modelo original nunca é alterado: ele é só lido, e o resultado é gravado no destino.
"""
import json, os, re, sys
from docx import Document
from docx.shared import Pt
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

CREDENCIAIS = os.path.normpath(os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    '..', '..', '..', '..', '..', '..', '..', '..', '5_Acervo', '2_Modelos Gerais', 'Credenciais'))
LINHA_DATA = '____ de ______________ de 20____'


def todas_tabelas(container):
    for t in container.tables:
        yield t
        for row in t.rows:
            for cell in row.cells:
                yield from todas_tabelas(cell)


def todos_paragrafos(doc):
    yield from doc.paragraphs
    for t in todas_tabelas(doc):
        for row in t.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def sem_linha(cell):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tcPr = cell._element.get_or_add_tcPr()
    bd = tcPr.find(qn('w:tcBorders'))
    if bd is not None:
        for e in bd.findall(qn('w:bottom')):
            bd.remove(e)
        nil = OxmlElement('w:bottom'); nil.set(qn('w:val'), 'nil'); bd.append(nil)


def preencher_campos(doc, campos):
    usados, faltou = set(), []
    linhas = []
    for t in todas_tabelas(doc):
        for row in t.rows:
            cells = row.cells
            if len(cells) == 2 and not cells[1].tables:
                linhas.append(cells)
    for rotulo, valor in campos:
        alvo = None
        for i, (a, b) in enumerate(linhas):
            if i in usados:
                continue
            if a.text.strip().lower() == rotulo.strip().lower() and not b.text.strip():
                alvo = i; break
        if alvo is None:
            faltou.append(rotulo); continue
        usados.add(alvo)
        a, b = linhas[alvo]
        r = b.paragraphs[0].add_run(valor); r.font.size = Pt(8.5)
        # preenchido: rótulo e valor alinhados pelo topo (endereço longo quebra em 2 linhas) e sem a
        # linha pontilhada de preenchimento à mão, que só faz sentido no campo vazio
        a.vertical_alignment = b.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        sem_linha(b)
        # sem espaço de escrita à mão, a linha pode ser mais baixa (cresce sozinha se o texto quebrar)
        tr = b._tc.getparent()
        from docx.table import _Row
        from docx.shared import Cm
        _Row(tr, None).height = Cm(0.5)
    return faltou


def substituir(doc, trocas):
    pendentes = dict(trocas)
    for p in todos_paragrafos(doc):
        for chave in list(trocas):
            if chave not in p.text:
                continue
            feito = False
            for r in p.runs:
                if chave in r.text:
                    r.text = r.text.replace(chave, trocas[chave]); feito = True
            if not feito:
                # texto quebrado em vários runs: junta no primeiro, mantendo a formatação dele
                txt = p.text.replace(chave, trocas[chave])
                for r in p.runs[1:]:
                    r.text = ''
                p.runs[0].text = txt
            pendentes.pop(chave, None)
    return list(pendentes)


def remover_poderes_especificos(doc):
    body = doc.element.body
    for p in doc.paragraphs:
        if 'Poderes Específicos' in p.text:
            prox = p._p.getnext()
            body.remove(p._p)
            if prox is not None and prox.tag.endswith('}tbl'):
                body.remove(prox)
            return True
    return False


def expandir_lista(doc, itens_por_chave):
    import copy
    faltou = []
    for chave, itens in itens_por_chave.items():
        alvo = next((p for p in todos_paragrafos(doc) if p.text.strip().lstrip('•').strip().startswith(chave)
                     or chave in p.text), None)
        if alvo is None:
            faltou.append(chave); continue
        anterior = alvo._p
        for i, item in enumerate(itens):
            novo = copy.deepcopy(alvo._p) if i else alvo._p
            if i:
                anterior.addnext(novo); anterior = novo
            runs = novo.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r')
            # run 0 = marcador (•	); o texto vai todo no último run, os demais ficam vazios
            textos = [r.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t') for r in runs]
            for t in textos[1:-1]:
                if t is not None: t.text = ''
            textos[-1].text = item
    return faltou


def altura_quadro(doc, cm):
    from docx.shared import Cm
    for t in doc.tables:
        if t.rows and t.rows[0].cells[0].text.strip().upper() == 'OBJETO':
            for row in t.rows:
                row.height = Cm(cm)
            return True
    return False


def remover_linhas(doc, inicios):
    feitos = set()
    for p in list(todos_paragrafos(doc)):
        t = p.text.strip()
        for ini in inicios:
            if t.startswith(ini):
                p._p.getparent().remove(p._p); feitos.add(ini)
    return [i for i in inicios if i not in feitos]


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    if len(args) != 3:
        sys.exit(__doc__)
    modelo, destino, dados_path = args
    if not modelo.lower().endswith('.docx'):
        modelo += '.docx'
    if not os.path.isabs(modelo):
        modelo = os.path.join(CREDENCIAIS, modelo)
    if not os.path.exists(modelo):
        sys.exit(f'Modelo não encontrado: {modelo}\nDisponíveis: ' + ', '.join(sorted(os.listdir(CREDENCIAIS))))
    destino = os.path.abspath(destino)
    if os.path.normcase(os.path.dirname(destino)).startswith(os.path.normcase(CREDENCIAIS)):
        sys.exit('Destino dentro da pasta Credenciais: proibido. Os modelos ficam intactos; salve na pasta do cliente.')
    if os.path.exists(destino) and '--sobrescrever' not in sys.argv:
        sys.exit(f'Já existe um arquivo em {destino}. Use outro nome ou --sobrescrever (só com autorização da Paula).')
    dados = json.load(open(dados_path, encoding='utf-8'))

    doc = Document(modelo)
    avisos = []
    if dados.get('campos'):
        f = preencher_campos(doc, dados['campos'])
        if f: avisos.append('Rótulos não encontrados/sem espaço vazio: ' + ', '.join(f))
    trocas = dict(dados.get('substituir', {}))
    if dados.get('data'):
        trocas[LINHA_DATA] = dados['data']
    if trocas:
        f = substituir(doc, trocas)
        if f: avisos.append('Textos não encontrados no modelo: ' + ' | '.join(f))
    if dados.get('sem_poderes_especificos'):
        if not remover_poderes_especificos(doc):
            avisos.append('Quadro de Poderes Específicos não encontrado.')

    if dados.get('expandir_lista'):
        f = expandir_lista(doc, dados['expandir_lista'])
        if f: avisos.append('Itens de lista não encontrados: ' + ' | '.join(f))
    if dados.get('altura_quadro'):
        if not altura_quadro(doc, dados['altura_quadro']):
            avisos.append('Quadro Resumo não encontrado pra ajustar altura.')
    if dados.get('remover_linhas'):
        f = remover_linhas(doc, dados['remover_linhas'])
        if f: avisos.append('Linhas a remover não encontradas: ' + ', '.join(f))

    os.makedirs(os.path.dirname(destino), exist_ok=True)
    doc.save(destino)
    restantes = sorted({m for p in todos_paragrafos(doc) for m in re.findall(r'\[[^\]]{2,}\]', p.text)})
    print('Salvo em:', destino)
    for a in avisos:
        print('AVISO:', a)
    if restantes:
        print('Placeholders ainda no documento:', ' | '.join(restantes))


if __name__ == '__main__':
    main()
